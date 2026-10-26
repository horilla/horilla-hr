"""
horilla_middlewares.py

This module is used to register horilla's middlewares without affecting the horilla/settings.py
"""

import contextlib
import csv
import datetime
import io
import json
import os
import re
import secrets
import threading
import zipfile
from contextvars import ContextVar
from urllib.parse import unquote

from django.conf import settings
from django.contrib import messages
from django.core.cache import cache
from django.core.exceptions import ObjectDoesNotExist
from django.http import HttpResponse, HttpResponseNotAllowed, JsonResponse
from django.shortcuts import render
from django.template.loader import render_to_string
from django.utils.datastructures import MultiValueDictKeyError
from django.utils.http import content_disposition_header
from openpyxl import Workbook, load_workbook
from xhtml2pdf import pisa

from horilla.config import logger
from horilla.export_safety import safe_cell
from horilla.http.response import HorillaRedirect

_request_var = ContextVar("request", default=None)
current_company_id = ContextVar("current_company_id", default=None)
_thread_local_state = ContextVar("thread_local_state", default={})


class _ThreadLocalProxy:
    def __getattr__(self, name):
        if name == "request":
            return _request_var.get()
        state = _thread_local_state.get()
        if name in state:
            return state[name]
        raise AttributeError(name)

    def __setattr__(self, name, value):
        if name == "request":
            _request_var.set(value)
        else:
            state = dict(_thread_local_state.get())
            state[name] = value
            _thread_local_state.set(state)


_thread_locals = _ThreadLocalProxy()


class ThreadLocalMiddleware:
    """
    ThreadLocalMiddleWare
    """

    def __init__(self, get_response):
        self.get_response = get_response

    def __call__(self, request):
        _thread_locals.request = request
        try:
            return self.get_response(request)
        finally:
            # Clear on the way out. Threads are reused -- by a worker for the
            # next request, and by the test runner for the next test -- so a
            # request left here outlives the one it belongs to. Anything that
            # reads it later (HorillaModel.save() stamps created_by/modified_by
            # from request.user) then attributes work to a stale user. In tests
            # that surfaces as an IntegrityError at teardown: modified_by_id
            # points at a user whose transaction has already rolled back.
            _thread_locals.request = None


class DefaultLanguageMiddleware:
    """
    Runs before django.middleware.locale.LocaleMiddleware to stop it from
    picking a language via the browser's Accept-Language header. Django's
    set_language view (used by the navbar language switcher) only persists
    an explicit choice in the LANGUAGE_COOKIE_NAME cookie, so that cookie
    is the sole signal that should override settings.LANGUAGE_CODE.
    """

    def __init__(self, get_response):
        self.get_response = get_response

    def __call__(self, request):
        if settings.LANGUAGE_COOKIE_NAME not in request.COOKIES:
            request.META.pop("HTTP_ACCEPT_LANGUAGE", None)
        return self.get_response(request)


class MethodNotAllowedMiddleware:
    def __init__(self, get_response):
        self.get_response = get_response

    def __call__(self, request):
        response = self.get_response(request)
        if isinstance(response, HttpResponseNotAllowed):
            return render(request, "405.html", status=405)
        return response


class SVGSecurityMiddleware:
    def __init__(self, get_response):
        self.get_response = get_response

    def __call__(self, request):
        response = self.get_response(request)

        # Apply security headers to SVG files
        if request.path.endswith(".svg") and response.status_code == 200:
            response["Content-Security-Policy"] = (
                "default-src 'none'; style-src 'unsafe-inline';"
            )
            response["X-Content-Type-Options"] = "nosniff"

        return response


class MissingParameterMiddleware:
    def __init__(self, get_response):
        self.get_response = get_response

    def __call__(self, request):
        return self.get_response(request)

    def process_exception(self, request, exception):
        if isinstance(exception, KeyError):
            missing_key = str(exception).strip("'")
            message = f"Required parameter '{missing_key}' is missing from the request."

            logger.error(message)

            return HorillaRedirect(request, message=message)

        elif isinstance(exception, ObjectDoesNotExist):
            logger.error(f"{exception.__class__.__name__}: {exception}")

            return HorillaRedirect(
                request, message="The requested item could not be found."
            )

        return None


def set_selected_company(company_id):
    current_company_id.set(company_id)


def get_selected_company():
    return current_company_id.get()


class BulkDeleteSummaryMiddleware:
    """
    Drops the toast messages of a bulk delete the browser is already reporting.

    static/index/index.js wraps every bulk-delete call: progress dialog, then a
    Total / Deleted / Failed result dialog worked out from which ids still
    exist. It marks those requests with X-Hlv-Bulk-Delete. The views still queue
    a toast per record ("You cannot delete X"), which would show up on top of
    that dialog (or, after a page reload, long after it), so they're discarded
    here instead of being edited out of every view. Bulk approve / reject views
    do the same by returning X-Hlv-Bulk-Result
    (horilla.http.response.with_bulk_result).

    Must sit after MessageMiddleware in MIDDLEWARE: responses unwind in reverse,
    so this runs before the messages are written back to the session.
    """

    def __init__(self, get_response):
        self.get_response = get_response

    def __call__(self, request):
        response = self.get_response(request)
        if request.headers.get("X-Hlv-Bulk-Delete") or response.has_header(
            "X-Hlv-Bulk-Result"
        ):
            storage = getattr(request, "_messages", None)
            if storage is not None:
                # Iterating marks every queued message as used.
                list(storage)
        return response


class ExportDownloadCookieMiddleware:
    """
    Tells the browser when an export download has started.

    Export forms submit natively, so the page itself never learns when the
    file arrives. static/index/index.js shows an "Exporting..." dialog for
    them and adds a random `hlv_export_token` to the request; when the
    response is a file, the token is echoed back in the `hlv_export_done`
    cookie, which is the dialog's cue to close.
    """

    FILE_TYPES = (
        "application/pdf",
        "application/zip",
        "application/octet-stream",
        "application/vnd.",
        "text/csv",
    )

    def __init__(self, get_response):
        self.get_response = get_response

    def __call__(self, request):
        response = self.get_response(request)
        token = request.GET.get("hlv_export_token") or request.POST.get(
            "hlv_export_token"
        )
        if token and self._is_file(response):
            response.set_cookie(
                "hlv_export_done", token[:64], max_age=120, samesite="Lax"
            )
        return response

    def _is_file(self, response):
        if response.has_header("Content-Disposition"):
            return True
        content_type = response.get("Content-Type", "")
        return content_type.startswith(self.FILE_TYPES)


class ExportOptionsMiddleware:
    """
    Applies the file name and format picked in an export modal to the file the
    export view produced.

    Dozens of views each build their own workbook under a fixed name.
    static/index/index.js adds a "File name" and a "Format" field to every
    export form (`hlv_export_name`, `hlv_export_format`); rather than teaching
    every view about them, the resulting spreadsheet is renamed and, when
    another format was asked for, converted here. xlsx and csv exports can
    become xlsx, csv (a zip when the workbook has several sheets) or pdf; any
    other kind of file (json, ...) is only renamed.
    """

    FORMATS = ("xlsx", "csv", "pdf")
    XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def __init__(self, get_response):
        self.get_response = get_response

    CACHE_SECONDS = 300
    CACHE_MAX_BYTES = 20 * 1024 * 1024

    def __call__(self, request):
        response = self._cached_response(request) or self.get_response(request)

        # Step one of a quick (ajax) export: tell the browser which columns the
        # file has, so it can offer the column picker, and keep the file for
        # step two so the export view isn't run twice.
        if self._param(request, "hlv_export_probe") and self._is_workbook(response):
            return self._probe(request, response)

        requested_format = self._param(request, "hlv_export_format").lower()
        requested_name = self._clean_name(self._param(request, "hlv_export_name"))
        columns = self._columns(request)
        if requested_format not in self.FORMATS:
            requested_format = ""
        if not (requested_format or requested_name or columns):
            return response
        if getattr(response, "streaming", False) or response.status_code != 200:
            return response
        disposition = response.get("Content-Disposition", "")
        if not disposition.lower().startswith("attachment"):
            return response

        base, ext = self._split_filename(disposition)
        final_name = requested_name or base
        final_ext = ext
        wants_format = requested_format and requested_format != ext
        if ext in ("xlsx", "csv") and (wants_format or columns):
            target = requested_format or ext
            try:
                content, content_type, final_ext = self._convert(
                    request, response.content, ext, target, final_name, columns
                )
            except Exception:
                logger.exception("Could not convert export %s to %s", ext, target)
                final_ext = ext
            else:
                response.content = content
                response["Content-Type"] = content_type
                if response.has_header("Content-Length"):
                    response["Content-Length"] = str(len(content))
        response["Content-Disposition"] = content_disposition_header(
            True, f"{final_name}.{final_ext}"
        )
        return response

    # -- quick export, step one / step two ---------------------------------
    def _is_workbook(self, response):
        if getattr(response, "streaming", False) or response.status_code != 200:
            return False
        disposition = response.get("Content-Disposition", "")
        if not disposition.lower().startswith("attachment"):
            return False
        return self._split_filename(disposition)[1] == "xlsx"

    def _probe(self, request, response):
        base, _ext = self._split_filename(response["Content-Disposition"])
        try:
            rows = self._sheets(response.content, "xlsx")[0][1]
        except Exception:
            logger.exception("Could not read export workbook")
            return response
        header_at = self._header_index(rows)
        columns = self._header_labels(rows[header_at] if rows else [])
        token = secrets.token_urlsafe(16)
        user = getattr(request, "user", None)
        if (
            user is not None
            and user.is_authenticated
            and len(response.content) <= self.CACHE_MAX_BYTES
        ):
            with contextlib.suppress(Exception):
                cache.set(
                    f"hlv_export:{token}",
                    {
                        "user": user.pk,
                        "content": response.content,
                        "type": response.get("Content-Type", self.XLSX_TYPE),
                        "disposition": response["Content-Disposition"],
                    },
                    self.CACHE_SECONDS,
                )
        return JsonResponse({"token": token, "filename": base, "columns": columns})

    def _cached_response(self, request):
        """The file kept by the probe, when this request asks for it."""
        token = self._param(request, "hlv_export_use")
        user = getattr(request, "user", None)
        if not token or user is None or not user.is_authenticated:
            return None
        with contextlib.suppress(Exception):
            entry = cache.get(f"hlv_export:{token}")
            if entry and entry["user"] == user.pk:
                response = HttpResponse(entry["content"], content_type=entry["type"])
                response["Content-Disposition"] = entry["disposition"]
                return response
        return None

    def _columns(self, request):
        """Positions of the columns to keep, from `hlv_export_columns`."""
        raw = self._param(request, "hlv_export_columns")
        if not raw:
            return None
        try:
            positions = [int(i) for i in json.loads(raw)]
        except (ValueError, TypeError):
            return None
        return positions or None

    @staticmethod
    def _param(request, key):
        return (request.POST.get(key) or request.GET.get(key) or "").strip()

    @staticmethod
    def _clean_name(name):
        """A safe file name stem: no path, no extension, no odd characters."""
        name = re.sub(r"[^\w\s().-]", "", name, flags=re.UNICODE).strip(" .")
        for suffix in (".xlsx", ".csv", ".pdf"):
            if name.lower().endswith(suffix):
                name = name[: -len(suffix)]
        return name[:100].strip()

    @staticmethod
    def _split_filename(disposition):
        match = re.search(
            r"filename\*?=(?:UTF-8\'\')?\"?([^\";]+)\"?", disposition, re.I
        )
        filename = unquote(match.group(1)) if match else "export"
        base, dot, ext = filename.rpartition(".")
        return (base, ext.lower()) if dot else (filename, "")

    # -- reading ----------------------------------------------------------
    @staticmethod
    def _cell(value):
        if value is None:
            return ""
        if isinstance(value, datetime.datetime):
            return value.strftime("%Y-%m-%d %H:%M:%S" if value.time() else "%Y-%m-%d")
        if isinstance(value, datetime.date):
            return value.isoformat()
        return value

    def _sheets(self, content, ext):
        if ext == "csv":
            text = content.decode("utf-8-sig", errors="replace")
            sheets = [("Sheet1", list(csv.reader(io.StringIO(text))))]
        else:
            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=True
            )
            sheets = [
                (
                    ws.title,
                    [
                        [self._cell(c) for c in row]
                        for row in ws.iter_rows(values_only=True)
                    ],
                )
                for ws in workbook.worksheets
            ]
        cleaned = []
        for title, rows in sheets:
            rows = [row for row in rows if any(str(c).strip() for c in row)]
            # Workbooks often carry empty columns past the data; drop them.
            width = max(
                (
                    max((i + 1 for i, c in enumerate(row) if str(c).strip()), default=0)
                    for row in rows
                ),
                default=0,
            )
            cleaned.append((title, [list(row[:width]) for row in rows]))
        return cleaned

    # -- writing ----------------------------------------------------------
    def _convert(self, request, content, ext, target, name, columns=None):
        sheets = self._sheets(content, ext)
        if columns:
            first_title, first_rows = sheets[0]
            sheets[0] = (
                first_title,
                [[row[i] for i in columns if i < len(row)] for row in first_rows],
            )
        if target == "xlsx":
            workbook = Workbook()
            workbook.remove(workbook.active)
            for title, rows in sheets:
                ws = workbook.create_sheet(title[:31] or "Sheet")
                for row in rows:
                    ws.append([safe_cell(c) for c in row])
            buffer = io.BytesIO()
            workbook.save(buffer)
            return buffer.getvalue(), self.XLSX_TYPE, "xlsx"
        if target == "csv":
            if len(sheets) == 1:
                return self._csv_bytes(sheets[0][1]), "text/csv; charset=utf-8", "csv"
            buffer = io.BytesIO()
            with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
                for title, rows in sheets:
                    archive.writestr(f"{name}-{title}.csv", self._csv_bytes(rows))
            return buffer.getvalue(), "application/zip", "zip"
        return self._pdf_bytes(request, sheets[0][1], name), "application/pdf", "pdf"

    @staticmethod
    def _csv_bytes(rows):
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        for row in rows:
            writer.writerow([safe_cell(c) for c in row])
        return buffer.getvalue().encode("utf-8-sig")

    @staticmethod
    def _header_index(rows):
        """The header is the first of the top rows that uses the full width
        (workbooks often start with a company / title banner)."""
        widths = [sum(1 for c in row if str(c).strip()) for row in rows[:10]]
        return widths.index(max(widths)) if widths else 0

    @staticmethod
    def _header_labels(row):
        labels, seen = [], {}
        for index, cell in enumerate(row):
            label = str(cell).strip() or f"Column {index + 1}"
            seen[label] = seen.get(label, 0) + 1
            labels.append(label if seen[label] == 1 else f"{label} ({seen[label]})")
        return labels

    def _pdf_bytes(self, request, rows, name):
        header_at = self._header_index(rows)
        headers = self._header_labels(rows[header_at] if rows else [])
        body = [
            {h: ("" if v is None else str(v)) for h, v in zip(headers, row)}
            for row in rows[header_at + 1 :]
        ]
        company = getattr(request, "selected_company_instance", None)
        logo_path = None
        if company and company.icon:
            with contextlib.suppress(Exception):
                if os.path.exists(company.icon.path):
                    logo_path = company.icon.path
        html = render_to_string(
            "generic/export_pdf.html",
            {
                "headers": headers,
                "rows": body,
                "company_name": (company.company if company else "") or "",
                "date_range": "",
                "report_title": name,
                "logo_path": logo_path,
                "landscape": len(headers) > 5,
            },
        )
        buffer = io.BytesIO()
        pisa.CreatePDF(html, dest=buffer)
        return buffer.getvalue()
