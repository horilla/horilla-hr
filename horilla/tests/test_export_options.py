"""
ExportOptionsMiddleware renames an export and converts it to the format picked
in the export modal (static/index/index.js adds `hlv_export_name` /
`hlv_export_format` to every export form).
"""

import csv
import io
import json
import zipfile
from unittest import mock

from django.http import HttpResponse
from django.test import RequestFactory, SimpleTestCase
from openpyxl import Workbook, load_workbook

from horilla.horilla_middlewares import ExportOptionsMiddleware

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _workbook(*sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets:
        ws = workbook.create_sheet(title)
        for row in rows:
            ws.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _xlsx_response(*sheets, filename="employee_export.xlsx"):
    response = HttpResponse(_workbook(*sheets), content_type=XLSX_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _run(response, user=None, **params):
    request = RequestFactory().post("/employee/work-info-export/", params)
    if user is not None:
        request.user = user
    return ExportOptionsMiddleware(lambda request: response)(request)


def _user(pk=7):
    return mock.Mock(is_authenticated=True, pk=pk)


ROWS = [["Name", "Email"], ["Adam", "adam@example.com"], ["Eve", "@eve"]]


class ExportOptionsMiddlewareTests(SimpleTestCase):
    def test_no_options_leaves_the_response_alone(self):
        response = _xlsx_response(("Sheet1", ROWS))
        result = _run(response)
        self.assertIn("employee_export.xlsx", result["Content-Disposition"])
        self.assertEqual(result["Content-Type"], XLSX_TYPE)

    def test_rename_only(self):
        result = _run(_xlsx_response(("Sheet1", ROWS)), hlv_export_name="April staff")
        self.assertIn('filename="April staff.xlsx"', result["Content-Disposition"])
        self.assertEqual(result["Content-Type"], XLSX_TYPE)

    def test_name_is_sanitised(self):
        result = _run(
            _xlsx_response(("Sheet1", ROWS)), hlv_export_name="../../etc/passwd.xlsx<>"
        )
        disposition = result["Content-Disposition"]
        self.assertNotIn("/", disposition)
        self.assertNotIn("<", disposition)
        self.assertIn("etcpasswd.xlsx", disposition.replace('"', ""))

    def test_xlsx_to_csv(self):
        result = _run(
            _xlsx_response(("Sheet1", ROWS)),
            hlv_export_name="staff",
            hlv_export_format="csv",
        )
        self.assertIn('filename="staff.csv"', result["Content-Disposition"])
        self.assertTrue(result["Content-Type"].startswith("text/csv"))
        rows = list(csv.reader(io.StringIO(result.content.decode("utf-8-sig"))))
        self.assertEqual(rows[0], ["Name", "Email"])
        self.assertEqual(rows[1], ["Adam", "adam@example.com"])
        # a formula-looking value is neutralised
        self.assertEqual(rows[2][1], "'@eve")

    def test_empty_trailing_columns_are_dropped(self):
        result = _run(
            _xlsx_response(
                ("Sheet1", [["A", "B", None, None], ["1", "2", None, None]])
            ),
            hlv_export_format="csv",
        )
        rows = list(csv.reader(io.StringIO(result.content.decode("utf-8-sig"))))
        self.assertEqual(rows, [["A", "B"], ["1", "2"]])

    def test_multi_sheet_workbook_to_csv_is_a_zip(self):
        result = _run(
            _xlsx_response(("People", ROWS), ("Teams", [["Team"], ["A"]])),
            hlv_export_name="all",
            hlv_export_format="csv",
        )
        self.assertIn('filename="all.zip"', result["Content-Disposition"])
        with zipfile.ZipFile(io.BytesIO(result.content)) as archive:
            self.assertEqual(
                sorted(archive.namelist()), ["all-People.csv", "all-Teams.csv"]
            )

    @mock.patch(
        "horilla.horilla_middlewares.render_to_string",
        return_value="<html><body><table><tr><td>x</td></tr></table></body></html>",
    )
    def test_xlsx_to_pdf(self, render):
        result = _run(
            _xlsx_response(("Sheet1", [["Company"], *ROWS])),
            hlv_export_name="staff",
            hlv_export_format="pdf",
        )
        self.assertIn('filename="staff.pdf"', result["Content-Disposition"])
        self.assertEqual(result["Content-Type"], "application/pdf")
        self.assertTrue(result.content.startswith(b"%PDF"))
        # The banner row is skipped: the header is the widest of the top rows.
        context = render.call_args.args[1]
        self.assertEqual(context["headers"], ["Name", "Email"])
        self.assertEqual(
            context["rows"][0], {"Name": "Adam", "Email": "adam@example.com"}
        )

    def test_csv_to_xlsx(self):
        response = HttpResponse("Name,Email\nAdam,a@b.c\n", content_type="text/csv")
        response["Content-Disposition"] = 'attachment; filename="people.csv"'
        result = _run(response, hlv_export_format="xlsx")
        self.assertIn('filename="people.xlsx"', result["Content-Disposition"])
        sheet = load_workbook(io.BytesIO(result.content)).active
        self.assertEqual([c.value for c in sheet[1]], ["Name", "Email"])

    def test_other_file_types_are_only_renamed(self):
        response = HttpResponse("{}", content_type="application/json")
        response["Content-Disposition"] = 'attachment; filename="data.json"'
        result = _run(response, hlv_export_name="dump", hlv_export_format="pdf")
        self.assertIn('filename="dump.json"', result["Content-Disposition"])
        self.assertEqual(result["Content-Type"], "application/json")

    def test_non_attachment_responses_are_ignored(self):
        result = _run(
            HttpResponse("<p>hi</p>"), hlv_export_name="x", hlv_export_format="csv"
        )
        self.assertFalse(result.has_header("Content-Disposition"))


class QuickExportFlowTests(SimpleTestCase):
    """Probe (which columns?) then use (the file kept by the probe)."""

    BANNER_ROWS = [
        ["Company banner"],
        ["Name", "Email", "City"],
        ["Adam", "a@b.c", "Pune"],
    ]

    def test_probe_lists_the_columns_and_keeps_the_file(self):
        result = _run(
            _xlsx_response(("Sheet1", self.BANNER_ROWS)),
            user=_user(),
            hlv_export_probe="1",
        )
        payload = json.loads(result.content)
        self.assertEqual(payload["columns"], ["Name", "Email", "City"])
        self.assertEqual(payload["filename"], "employee_export")
        self.assertTrue(payload["token"])

    def test_probe_leaves_non_workbooks_alone(self):
        response = HttpResponse("<p>hi</p>")
        result = _run(response, user=_user(), hlv_export_probe="1")
        self.assertEqual(result.content, b"<p>hi</p>")

    def test_use_serves_the_kept_file_with_the_chosen_columns(self):
        probe = _run(
            _xlsx_response(
                ("Sheet1", [["Name", "Email", "City"], ["Adam", "a@b.c", "Pune"]])
            ),
            user=_user(),
            hlv_export_probe="1",
        )
        token = json.loads(probe.content)["token"]

        def never_called(request):
            raise AssertionError("the export view should not run again")

        request = RequestFactory().get(
            "/anything/",
            {
                "hlv_export_use": token,
                "hlv_export_columns": "[0, 2]",
                "hlv_export_format": "csv",
                "hlv_export_name": "People",
            },
        )
        request.user = _user()
        result = ExportOptionsMiddleware(never_called)(request)
        self.assertIn('filename="People.csv"', result["Content-Disposition"])
        rows = list(csv.reader(io.StringIO(result.content.decode("utf-8-sig"))))
        self.assertEqual(rows, [["Name", "City"], ["Adam", "Pune"]])

    def test_another_users_token_is_ignored(self):
        probe = _run(
            _xlsx_response(("Sheet1", ROWS)), user=_user(7), hlv_export_probe="1"
        )
        token = json.loads(probe.content)["token"]
        fresh = _xlsx_response(("Sheet1", [["Other"], ["file"]]))
        request = RequestFactory().get("/x/", {"hlv_export_use": token})
        request.user = _user(8)
        result = ExportOptionsMiddleware(lambda request: fresh)(request)
        self.assertIs(result, fresh)

    def test_columns_only_keeps_xlsx(self):
        result = _run(
            _xlsx_response(("Sheet1", [["A", "B", "C"], ["1", "2", "3"]])),
            hlv_export_columns="[0, 2]",
        )
        self.assertIn("employee_export.xlsx", result["Content-Disposition"])
        sheet = load_workbook(io.BytesIO(result.content)).active
        self.assertEqual(
            [[c.value for c in row] for row in sheet.iter_rows()],
            [["A", "C"], ["1", "3"]],
        )
