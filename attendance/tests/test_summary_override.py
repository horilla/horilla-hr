"""
Regularising a month by stating its totals.

What HR knows at month end is "this person worked nineteen days, not three".
That is a statement about the month, not about any particular date, and turning
it into per-day decisions produces a record that was invented.

So the sheet is one row per employee and the stated totals replace the counted
ones. The three things that matter:

  * a blank count means "leave that one as counted" — correcting only Absent is
    a normal thing to want, and blanking the rest would silently zero them;
  * a stated month settles its own conflicts, or payroll stays blocked on days
    the override was written to answer;
  * the period comes from the request, not from the sheet — otherwise a
    corrected March sheet could overwrite April.
"""

import datetime
from io import BytesIO

import pandas as pd
from django.test import TestCase
from django.urls import reverse

from attendance.models import AttendanceSummaryOverride
from attendance.views.summary import build_monthly_summary
from employee.models import Employee
from horilla.testkit import make_company, make_employee, make_user

FROM_DATE = datetime.date(2026, 4, 1)
TO_DATE = datetime.date(2026, 4, 30)

COLUMNS = ("Badge ID", "Present", "Paid Leave", "Unpaid Leave", "Absent", "Reason")


def sheet(rows, columns=COLUMNS):
    """An uploaded workbook, in memory."""
    buffer = BytesIO()
    pd.DataFrame(rows, columns=list(columns)).to_excel(buffer, index=False)
    buffer.seek(0)
    buffer.name = "summary.xlsx"
    return buffer


class Fixture(TestCase):
    def setUp(self):
        self.user = make_user("sumadmin", is_superuser=True)
        self.company = make_company("Summary Co")
        make_employee(
            company=self.company, email="sumadmin@test.horilla", user=self.user
        )
        self.client.force_login(self.user)

        self.employee = make_employee(
            company=self.company, email="worker@test.horilla", first_name="Worker"
        )
        # The factory does not set one, and badge is what the sheet matches on.
        self.employee.badge_id = "TST001"
        self.employee.save()
        self.badge = self.employee.badge_id

    def upload(self, payload, **extra):
        data = {
            "regularisation_file": payload,
            "from_date": FROM_DATE.isoformat(),
            "to_date": TO_DATE.isoformat(),
        }
        data.update(extra)
        return self.client.post(reverse("attendance-regularisation-import"), data)

    def override(self):
        return AttendanceSummaryOverride.objects.filter(
            employee_id=self.employee
        ).first()

    def summary_row(self):
        rows, _working, _totals = build_monthly_summary(
            FROM_DATE, TO_DATE, Employee.objects.filter(pk=self.employee.pk)
        )
        return rows[0]

    def reconciling(self, present=0, paid_leave=0, unpaid_leave=0, **extra):
        """
        A sheet row that accounts for every day of the period.

        Absent absorbs the remainder, so a test can state the figures it cares
        about without every one of them having to restate a whole month. The
        working-day total comes from the engine, so this reconciles exactly the
        way an untouched download does.
        """
        counted = self.summary_row()
        working = (
            counted["present"]
            + counted["paid_leave"]
            + counted["unpaid_leave"]
            + counted["absent"]
        )
        row = {
            "Badge ID": self.badge,
            "Present": present,
            "Paid Leave": paid_leave,
            "Unpaid Leave": unpaid_leave,
            "Absent": working - present - paid_leave - unpaid_leave,
        }
        row.update(extra)
        return sheet([row])


class TemplateTests(Fixture):
    def download(self):
        return self.client.get(
            reverse("attendance-regularisation-template"),
            {"from_date": FROM_DATE, "to_date": TO_DATE},
        )

    def test_it_downloads_as_a_workbook(self):
        response = self.download()
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn("attachment;", response["Content-Disposition"])

    def test_it_is_one_row_per_employee_not_per_day(self):
        """The whole point of the change: nobody wants to see days."""
        frame = pd.read_excel(BytesIO(self.download().content), sheet_name="Summary")
        self.assertNotIn("Date", frame.columns)
        self.assertEqual(len(frame[frame["Badge ID"] == self.badge]), 1)

    def test_it_comes_pre_filled_with_what_was_counted(self):
        """
        Empty would mean reconstructing the whole month. The counted figures
        are mostly right; only the wrong ones need touching.
        """
        frame = pd.read_excel(BytesIO(self.download().content), sheet_name="Summary")
        row = frame[frame["Badge ID"] == self.badge].iloc[0]

        for column in ("Present", "Paid Leave", "Unpaid Leave", "Absent"):
            self.assertFalse(pd.isna(row[column]), f"{column} came back blank")

    def test_it_says_which_columns_are_read_back(self):
        notes = pd.read_excel(
            BytesIO(self.download().content), sheet_name="How to fill this in"
        )
        editable = dict(zip(notes["Column"], notes["Editable"]))

        self.assertEqual(editable["Present"], "yes")
        self.assertEqual(editable["Absent"], "yes")
        # Roster and calendar, not anybody's judgement.
        self.assertEqual(editable["Week Off"], "no")
        self.assertEqual(editable["Holiday"], "no")


class TotalColumnTests(Fixture):
    def download(self):
        return self.client.get(
            reverse("attendance-regularisation-template"),
            {"from_date": FROM_DATE, "to_date": TO_DATE},
        )

    def test_the_sheet_carries_a_total_column(self):
        frame = pd.read_excel(BytesIO(self.download().content), sheet_name="Summary")
        self.assertIn("Total", frame.columns)

    def test_it_sits_beside_the_columns_it_sums(self):
        """A total three columns away from its own addends is one nobody checks."""
        frame = pd.read_excel(BytesIO(self.download().content), sheet_name="Summary")
        columns = list(frame.columns)
        self.assertEqual(columns[columns.index("Holiday") + 1], "Total")

    def test_it_is_a_live_formula_not_a_number(self):
        """
        A total computed at download time is stale the moment anybody types,
        and a stale total is worse than none -- it looks authoritative.
        """
        import openpyxl

        book = openpyxl.load_workbook(BytesIO(self.download().content))
        summary = book["Summary"]
        header = [cell.value for cell in summary[1]]
        column = header.index("Total") + 1

        formula = summary.cell(row=2, column=column).value
        self.assertTrue(
            str(formula).startswith("=SUM("), f"expected a formula, got {formula!r}"
        )

    def test_the_notes_explain_it(self):
        notes = pd.read_excel(
            BytesIO(self.download().content), sheet_name="How to fill this in"
        )
        row = notes[notes["Column"] == "Total"].iloc[0]
        self.assertEqual(row["Editable"], "no")
        self.assertIn("red", str(row["Notes"]))


class ReconcilesTests(Fixture):
    """
    A row has to add up against itself.

    total_working comes from get_working_days(), which is called without an
    employee -- so it is the company default roster. Every other column on the
    row is that employee's own. On anybody whose roster differs, the sheet
    showed 23 working days beside 24 days of present/leave/absent, and nothing
    on it explained the extra day.
    """

    def test_working_days_match_the_columns_beside_them(self):
        row = self.summary_row()
        self.assertEqual(
            row["working_days"],
            row["present"] + row["paid_leave"] + row["unpaid_leave"] + row["absent"],
        )

    def test_the_sheet_carries_the_reconciling_figure(self):
        response = self.client.get(
            reverse("attendance-regularisation-template"),
            {"from_date": FROM_DATE, "to_date": TO_DATE},
        )
        frame = pd.read_excel(BytesIO(response.content), sheet_name="Summary")
        row = frame[frame["Badge ID"] == self.badge].iloc[0]

        self.assertEqual(
            row["Working Days"],
            row["Present"] + row["Paid Leave"] + row["Unpaid Leave"] + row["Absent"],
        )

    def test_a_stated_month_recomputes_it(self):
        """It is derived, so an override that changed the counts has to move it."""
        AttendanceSummaryOverride.objects.create(
            employee_id=self.employee,
            from_date=FROM_DATE,
            to_date=TO_DATE,
            present=19,
            absent=0,
        )
        row = self.summary_row()
        self.assertEqual(
            row["working_days"],
            19 + row["paid_leave"] + row["unpaid_leave"] + 0,
        )


class SheetUsabilityTests(Fixture):
    def workbook(self):
        import openpyxl

        response = self.client.get(
            reverse("attendance-regularisation-template"),
            {"from_date": FROM_DATE, "to_date": TO_DATE},
        )
        return openpyxl.load_workbook(BytesIO(response.content))

    def test_only_the_header_row_is_frozen(self):
        """
        Freezing Badge ID and Employee too slid Present -- the first editable
        column -- behind the frozen pane as soon as anybody scrolled right.
        """
        summary = self.workbook()["Summary"]
        self.assertEqual(summary.freeze_panes, "A2")


class ImportTests(Fixture):
    def test_stated_totals_are_stored(self):
        self.upload(
            self.reconciling(
                present=19, paid_leave=2, Reason="Clocking machine was down"
            )
        )

        stored = self.override()
        self.assertEqual(stored.present, 19)
        self.assertEqual(stored.paid_leave, 2)
        self.assertEqual(stored.note, "Clocking machine was down")

    def test_a_blank_count_is_left_as_counted(self):
        """
        Correcting only Absent is normal. Blanking the rest would silently zero
        a month nobody said anything about.
        """
        counted = self.summary_row()
        self.upload(
            sheet(
                [
                    {
                        "Badge ID": self.badge,
                        "Present": "",
                        "Paid Leave": "",
                        "Unpaid Leave": "",
                        "Absent": counted["absent"],
                    }
                ]
            )
        )

        stored = self.override()
        self.assertIsNone(stored.present)
        self.assertIsNone(stored.paid_leave)
        self.assertEqual(stored.absent, counted["absent"])

    def test_a_row_with_nothing_stated_writes_nothing(self):
        self.upload(
            sheet(
                [
                    {
                        "Badge ID": self.badge,
                        "Present": "",
                        "Paid Leave": "",
                        "Unpaid Leave": "",
                        "Absent": "",
                    }
                ]
            )
        )
        self.assertIsNone(self.override())

    def test_re_uploading_updates_rather_than_duplicating(self):
        for present in (19, 20):
            self.upload(self.reconciling(present=present))

        self.assertEqual(
            AttendanceSummaryOverride.objects.filter(employee_id=self.employee).count(),
            1,
        )
        self.assertEqual(self.override().present, 20)

    def test_the_period_comes_from_the_request_not_the_sheet(self):
        """
        A column the sheet could edit would let a corrected March upload
        overwrite April.
        """
        self.upload(self.reconciling(present=19))

        stored = self.override()
        self.assertEqual(stored.from_date, FROM_DATE)
        self.assertEqual(stored.to_date, TO_DATE)

    def test_without_a_period_nothing_is_written(self):
        response = self.client.post(
            reverse("attendance-regularisation-import"),
            {"regularisation_file": self.reconciling(present=19)},
        )
        self.assertEqual(response.status_code, 302)
        self.assertIsNone(self.override())

    def test_it_returns_to_where_it_was_started_from(self):
        response = self.upload(
            self.reconciling(present=19),
            next=reverse("payroll-batch-review"),
        )
        self.assertEqual(response["Location"], reverse("payroll-batch-review"))

    def test_an_off_site_next_is_refused(self):
        """An unchecked next is an open redirect."""
        response = self.upload(
            self.reconciling(present=19),
            next="https://example.invalid/steal",
        )
        self.assertEqual(response["Location"], reverse("attendance-monthly-summary"))


class AppliedToTheSummaryTests(Fixture):
    def state(self, **counts):
        AttendanceSummaryOverride.objects.create(
            employee_id=self.employee,
            from_date=FROM_DATE,
            to_date=TO_DATE,
            **counts,
        )

    def test_the_summary_reports_what_was_stated(self):
        self.state(present=19, absent=0)
        row = self.summary_row()

        self.assertEqual(row["present"], 19)
        self.assertEqual(row["absent"], 0)
        self.assertTrue(row["summary_overridden"])

    def test_paid_days_follows_the_stated_figures(self):
        """
        It is derived, so leaving it computed would have the summary disagree
        with itself -- and payroll reads paid_days.
        """
        self.state(present=19, paid_leave=1)
        row = self.summary_row()

        self.assertEqual(
            row["paid_days"],
            19 + 1 + row["holiday"] + row["week_off"],
        )

    def test_a_stated_month_settles_its_own_conflicts(self):
        """
        Otherwise payroll stays blocked on exactly the days the override was
        written to answer.
        """
        self.state(present=19)
        row = self.summary_row()

        self.assertEqual(row["unresolved_conflicts"], 0)
        self.assertEqual(row["unresolved_conflict_dates"], [])

    def test_a_partial_statement_leaves_the_rest_counted(self):
        counted = self.summary_row()["present"]
        self.state(absent=0)

        self.assertEqual(self.summary_row()["present"], counted)
        self.assertEqual(self.summary_row()["absent"], 0)

    def test_another_period_is_untouched(self):
        self.state(present=19)
        rows, _w, _t = build_monthly_summary(
            datetime.date(2026, 5, 1),
            datetime.date(2026, 5, 31),
            Employee.objects.filter(pk=self.employee.pk),
        )
        self.assertNotIn("summary_overridden", rows[0])


class ConfirmationTests(Fixture):
    """
    Stating a month settles its conflicts. That is the point of it, and it is
    also a quiet thing to do to somebody's attendance record -- so when an
    upload would settle any, it stops and shows which.
    """

    def setUp(self):
        super().setUp()
        self.conflict_day = datetime.date(2026, 4, 15)

    def give_conflict(self):
        """
        A day carrying both attendance and approved leave, which is what
        build_monthly_summary counts as a conflict.
        """
        from attendance.models import Attendance
        from leave.models import LeaveRequest, LeaveType

        leave_type = LeaveType.objects.create(name="Annual", payment_type="paid")
        LeaveRequest.objects.create(
            employee_id=self.employee,
            leave_type_id=leave_type,
            start_date=self.conflict_day,
            end_date=self.conflict_day,
            status="approved",
        )
        Attendance.objects.create(
            employee_id=self.employee,
            attendance_date=self.conflict_day,
            attendance_clock_in_date=self.conflict_day,
            attendance_clock_in=datetime.time(9, 0),
            attendance_validated=True,
        )

    def test_without_conflicts_it_saves_straight_away(self):
        response = self.upload(self.reconciling(present=19))

        self.assertEqual(response.status_code, 302)
        self.assertNotIn("regularise=1", response["Location"])
        self.assertIsNotNone(self.override())

    def test_with_conflicts_nothing_is_written_yet(self):
        """
        The upload is validated and held. Somebody has to say yes before an
        attendance record changes.
        """
        self.give_conflict()
        response = self.upload(self.reconciling(present=19))

        self.assertIn("regularise=1", response["Location"])
        self.assertIsNone(self.override())

    def test_the_held_upload_survives_the_redirect(self):
        """
        The file itself cannot be carried across it, and asking for it again
        would mean re-picking it in the file dialog.
        """
        self.give_conflict()
        self.upload(self.reconciling(present=19))

        from attendance.views.regularisation_import import SESSION_KEY

        self.assertIn(SESSION_KEY, self.client.session)

    def test_the_question_names_the_days_it_would_settle(self):
        self.give_conflict()
        self.upload(self.reconciling(present=19))

        response = self.client.get(
            reverse("attendance-regularisation-confirm"), HTTP_HX_REQUEST="true"
        )
        self.assertContains(response, "Regularise these conflicts?")
        # The template renders each date as "j M", so 15 April reads "15 Apr".
        # The count and the dates both matter: "1 conflict" without saying
        # which day is not enough to decide on.
        self.assertContains(response, "15 Apr")
        self.assertContains(response, str(self.employee))

    def test_confirming_saves_it(self):
        self.give_conflict()
        self.upload(self.reconciling(present=19))

        self.client.post(reverse("attendance-regularisation-confirm"))
        self.assertIsNotNone(self.override())

    def test_cancelling_leaves_the_record_alone(self):
        """
        Nothing was written when the file was uploaded, so cancelling is not a
        rollback -- there is nothing to roll back.
        """
        self.give_conflict()
        self.upload(self.reconciling(present=19))

        self.client.post(reverse("attendance-regularisation-confirm"), {"cancel": "1"})
        self.assertIsNone(self.override())

    def test_the_held_upload_is_discarded_either_way(self):
        from attendance.views.regularisation_import import SESSION_KEY

        self.give_conflict()
        self.upload(self.reconciling(present=19))
        self.client.post(reverse("attendance-regularisation-confirm"), {"cancel": "1"})
        self.assertNotIn(SESSION_KEY, self.client.session)

    def test_confirming_with_nothing_held_does_nothing(self):
        response = self.client.get(
            reverse("attendance-regularisation-confirm"), HTTP_HX_REQUEST="true"
        )
        self.assertEqual(response.content, b"")


class RejectionTests(Fixture):
    def test_an_unknown_badge_is_reported_not_applied(self):
        response = self.upload(
            sheet(
                [{"Badge ID": "NOPE999", "Present": 19}],
                columns=("Badge ID", "Present"),
            )
        )

        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn("summary_errors", response["Content-Disposition"])
        self.assertIsNone(self.override())

    def test_a_count_that_is_not_a_number_is_reported(self):
        response = self.upload(
            sheet(
                [{"Badge ID": self.badge, "Present": "loads"}],
                columns=("Badge ID", "Present"),
            )
        )
        errors = pd.read_excel(BytesIO(response.content))
        self.assertIn("Present", str(errors["Problem"].iloc[0]))

    def test_a_negative_count_is_reported(self):
        response = self.upload(
            sheet(
                [{"Badge ID": self.badge, "Absent": -3}], columns=("Badge ID", "Absent")
            )
        )
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIsNone(self.override())

    def test_the_error_sheet_says_which_row(self):
        """Row 2 is the first data row -- row 1 is the header."""
        response = self.upload(
            sheet(
                [{"Badge ID": "NOPE999", "Present": 19}],
                columns=("Badge ID", "Present"),
            )
        )
        errors = pd.read_excel(BytesIO(response.content))
        self.assertEqual(int(errors["Row"].iloc[0]), 2)

    def test_one_bad_row_stops_the_whole_sheet(self):
        """
        Applied in part and then reconciled by hand is worse than rejected
        outright: the sheet is the record of what was stated, and it has to
        mean the same thing after uploading as before.
        """
        self.upload(
            sheet(
                [
                    {"Badge ID": self.badge, "Present": 19},
                    {"Badge ID": "NOPE999", "Present": 19},
                ],
                columns=("Badge ID", "Present"),
            )
        )
        self.assertIsNone(self.override())

    def test_a_sheet_missing_a_count_column_is_refused(self):
        response = self.upload(sheet([{"Badge ID": self.badge}], columns=("Badge ID",)))
        self.assertEqual(response.status_code, 302)
        self.assertIsNone(self.override())

    def test_a_month_that_accounts_for_every_day_is_allowed(self):
        self.upload(self.reconciling())
        self.assertIsNotNone(self.override())

    def test_more_days_than_the_period_has_is_refused(self):
        """
        The red in the spreadsheet is a hint, and a hint does not stop anybody
        uploading the file anyway.
        """
        response = self.upload(self.reconciling(present=25, Absent=25))

        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn(
            "30", str(pd.read_excel(BytesIO(response.content))["Problem"].iloc[0])
        )
        self.assertIsNone(self.override())

    def test_fewer_days_than_the_period_has_is_refused(self):
        """
        Days nobody accounted for. Payroll would compute from a month with
        holes in it, and only flagging the over case left this looking fine.
        """
        response = self.upload(self.reconciling(present=1, Absent=0))

        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIsNone(self.override())

    def test_a_blank_count_is_filled_from_what_was_counted(self):
        """
        Otherwise a partly-filled row -- the normal case -- would look short by
        whatever was left blank and be refused for it.
        """
        counted = self.summary_row()
        self.upload(
            sheet(
                [
                    {
                        "Badge ID": self.badge,
                        "Present": counted["present"],
                        "Paid Leave": "",
                        "Unpaid Leave": "",
                        "Absent": counted["absent"],
                    }
                ]
            )
        )
        self.assertIsNotNone(self.override())

    def test_a_get_does_nothing(self):
        self.client.get(reverse("attendance-regularisation-import"))
        self.assertIsNone(self.override())
