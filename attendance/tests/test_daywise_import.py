"""Importing by day: the Day-Wise grid as a sheet, a status dropdown in each cell, only changed cells applied."""

import datetime
import io
from datetime import time

from django.contrib.messages import get_messages
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from attendance.models import Attendance, AttendanceConflictResolution
from horilla.testkit import make_company, make_employee, make_user

DAY = datetime.date(2026, 9, 2)  # a Wednesday, in the past
HEADER_ROW = 2  # openpyxl is 1-based: weekdays, then the headers with the dates


class DayWiseImportTests(TestCase):
    def setUp(self):
        company = make_company("Day Wise Co")
        self.employee = make_employee(company=company, email="dw@test.horilla")
        user = make_user("dw-admin", is_superuser=True)
        make_employee(company=company, email="dw-admin@test.horilla", user=user)
        self.client.force_login(user)
        session = self.client.session
        session["selected_company"] = "all"
        session.save()
        Attendance.objects.create(
            employee_id=self.employee,
            attendance_date=DAY,
            attendance_clock_in_date=DAY,
            attendance_clock_in=time(9, 0),
            attendance_clock_out_date=DAY,
            attendance_clock_out=time(18, 0),
            attendance_worked_hour="09:00",
            minimum_hour="08:00",
            attendance_validated=True,
        )

    def download(self, start=DAY, end=DAY + datetime.timedelta(days=3)):
        return self.client.get(
            reverse("attendance-daywise-template"),
            {
                "from_date": start.isoformat(),
                "to_date": end.isoformat(),
                "search": self.employee.get_full_name(),
            },
        )

    def workbook(self, response):
        return load_workbook(io.BytesIO(response.content))

    def row_of(self, sheet):
        for row in range(HEADER_ROW + 1, sheet.max_row + 1):
            if sheet.cell(row, 1).value == self.employee.pk:
                return row
        self.fail("employee row not found")

    def upload(self, book):
        buffer = io.BytesIO()
        book.save(buffer)
        file = SimpleUploadedFile("sheet.xlsx", buffer.getvalue())
        return self.client.post(
            reverse("attendance-daywise-import"), {"regularisation_file": file}
        )

    def messages(self, response):
        return [str(m) for m in get_messages(response.wsgi_request)]

    def test_the_template_is_the_grid_with_a_header_row_per_date_and_a_status_in_each_cell(
        self,
    ):
        response = self.download()
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        book = self.workbook(response)
        self.assertEqual(book.sheetnames, ["Day-Wise", "How To Fill This In"])
        sheet = book["Day-Wise"]
        self.assertEqual(sheet.cell(1, 4).value, "Wed")
        self.assertEqual(
            [sheet.cell(HEADER_ROW, c).value for c in (1, 2, 3)],
            ["Employee ID", "Badge ID", "Employee"],
        )
        dates = [sheet.cell(HEADER_ROW, c).value.date() for c in range(4, 8)]
        self.assertEqual(dates[0], DAY)
        self.assertEqual(len(dates), 4)
        self.assertEqual(sheet.cell(self.row_of(sheet), 4).value, "P")

    def test_the_employee_column_has_the_name_only_not_the_badge_again(self):
        sheet = self.workbook(self.download())["Day-Wise"]
        row = self.row_of(sheet)
        self.assertEqual(sheet.cell(row, 3).value, self.employee.get_full_name())
        self.assertNotIn("(", sheet.cell(row, 3).value)

    def test_the_cells_have_a_status_dropdown_and_the_sheet_is_formatted(self):
        sheet = self.workbook(self.download())["Day-Wise"]
        validations = sheet.data_validations.dataValidation
        self.assertEqual(len(validations), 1)
        self.assertEqual(validations[0].type, "list")
        for code in ("P", "HD", "HA", "A", "PL", "UL", "WO", "H"):
            self.assertIn(code, validations[0].formula1)
        self.assertEqual(sheet.freeze_panes, "D3")
        self.assertTrue(sheet.column_dimensions["A"].hidden)
        self.assertTrue(sheet.conditional_formatting)
        self.assertEqual(
            sheet.cell(HEADER_ROW, 4).fill.fgColor.rgb[-6:].lower(), "f6b83c"
        )
        self.assertEqual(sheet.cell(HEADER_ROW, 4).number_format, "dd mmm")
        self.assertGreaterEqual(sheet.column_dimensions["D"].width, 8)

    def test_a_changed_cell_is_applied_as_a_regularisation(self):
        book = self.workbook(self.download())
        sheet = book["Day-Wise"]
        sheet.cell(self.row_of(sheet), 4).value = "A"
        response = self.upload(book)
        self.assertEqual(response.status_code, 302)
        resolution = AttendanceConflictResolution.objects.get(
            employee_id=self.employee, date=DAY
        )
        self.assertEqual(resolution.resolution, "absent")
        self.assertIn(
            "1 Days Regularised From The Day-Wise Sheet.", self.messages(response)
        )

    def test_an_untouched_sheet_changes_nothing(self):
        response = self.upload(self.workbook(self.download()))
        self.assertFalse(AttendanceConflictResolution.objects.exists())
        self.assertIn("No Changes Found In The Sheet.", self.messages(response))

    def test_the_last_upload_for_a_day_wins(self):
        for code, expected in (("HD", "half_present"), ("PL", "paid_leave")):
            book = self.workbook(self.download())
            sheet = book["Day-Wise"]
            sheet.cell(self.row_of(sheet), 4).value = code
            self.upload(book)
            self.assertEqual(
                AttendanceConflictResolution.objects.get(
                    employee_id=self.employee, date=DAY
                ).resolution,
                expected,
            )
        self.assertEqual(
            AttendanceConflictResolution.objects.filter(
                employee_id=self.employee, date=DAY
            ).count(),
            1,
        )

    def test_unknown_codes_and_upcoming_days_are_skipped_and_reported(self):
        tomorrow = datetime.date.today() + datetime.timedelta(days=1)
        book = self.workbook(self.download(DAY, tomorrow))
        sheet = book["Day-Wise"]
        row = self.row_of(sheet)
        sheet.cell(row, 5).value = "XYZ"
        sheet.cell(row, sheet.max_column).value = "P"
        response = self.upload(book)
        self.assertFalse(AttendanceConflictResolution.objects.exists())
        warning = [m for m in self.messages(response) if "Skipped" in m][0]
        self.assertIn("XYZ", warning)
        self.assertIn("Upcoming Day", warning)

    def test_a_file_that_is_not_a_day_wise_sheet_is_refused(self):
        from openpyxl import Workbook

        book = Workbook()
        book.active["A1"] = "nothing"
        response = self.upload(book)
        self.assertIn("Day-Wise Sheet", self.messages(response)[0])
        self.assertFalse(AttendanceConflictResolution.objects.exists())

    def test_a_range_longer_than_the_grid_allows_is_refused(self):
        response = self.download(DAY, DAY + datetime.timedelta(days=90))
        self.assertEqual(response.status_code, 302)

    def test_the_template_link_follows_the_screen_and_explains_what_it_covers(self):
        html = self.client.get(reverse("attendance-monthly-summary")).content.decode()
        self.assertIn('id="ms-template-link"', html)
        self.assertIn('id="ms-template-tip"', html)
        self.assertIn("function msSyncTemplateLink(params, selectedIds)", html)
        self.assertIn("msSyncTemplateLink(params, selectedIds);", html)
        self.assertIn("The Template Covers", html)

    def grid_cell(self):
        response = self.client.get(
            reverse("attendance-monthly-summary-table"),
            {
                "view": "daily",
                "from_date": DAY.isoformat(),
                "to_date": DAY.isoformat(),
                "search": self.employee.get_full_name(),
            },
            HTTP_HX_REQUEST="true",
        )
        row = [
            r for r in response.context["rows"] if r["employee"].pk == self.employee.pk
        ][0]
        return row["cells"][0]

    def import_code(self, code):
        book = self.workbook(self.download(DAY, DAY))
        sheet = book["Day-Wise"]
        sheet.cell(self.row_of(sheet), 4).value = code
        return self.upload(book)

    def test_an_imported_change_is_a_regularisation_the_grid_and_work_records_show(
        self,
    ):
        from attendance.models import WorkRecords

        self.import_code("A")
        cell = self.grid_cell()
        self.assertEqual(
            (cell["status"], cell["code"], cell["resolved"]), ("absent", "A", True)
        )
        record = WorkRecords.objects.get(employee_id=self.employee, date=DAY)
        self.assertEqual(
            (record.work_record_type, record.message), ("DFT", "Regularised: Absent")
        )

    def test_a_manual_regularisation_after_an_import_replaces_it(self):
        self.import_code("A")
        self.client.post(
            reverse("attendance-monthly-summary-conflict-resolve"),
            {
                "employee_id": self.employee.pk,
                "date": DAY.isoformat(),
                "resolution": "paid_leave",
                "from_date": DAY.isoformat(),
                "to_date": DAY.isoformat(),
            },
            HTTP_HX_REQUEST="true",
        )
        rows = AttendanceConflictResolution.objects.filter(
            employee_id=self.employee, date=DAY
        )
        self.assertEqual([r.resolution for r in rows], ["paid_leave"])
        self.assertEqual(self.grid_cell()["code"], "PL")

    def test_a_later_import_replaces_a_manual_regularisation(self):
        self.client.post(
            reverse("attendance-monthly-summary-conflict-resolve"),
            {
                "employee_id": self.employee.pk,
                "date": DAY.isoformat(),
                "resolution": "paid_leave",
                "from_date": DAY.isoformat(),
                "to_date": DAY.isoformat(),
            },
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(
            AttendanceConflictResolution.objects.get(
                employee_id=self.employee, date=DAY
            ).resolution,
            "paid_leave",
        )
        self.import_code("H")
        rows = AttendanceConflictResolution.objects.filter(
            employee_id=self.employee, date=DAY
        )
        self.assertEqual([r.resolution for r in rows], ["holiday"])
        self.assertEqual(self.grid_cell()["code"], "H")

    def test_the_summary_page_offers_the_day_wise_import_and_not_the_summary_sheet(
        self,
    ):
        html = self.client.get(reverse("attendance-monthly-summary")).content.decode()
        self.assertIn(reverse("attendance-daywise-import"), html)
        self.assertIn(reverse("attendance-daywise-template"), html)
        self.assertIn("Import Day-Wise Sheet", html)
        self.assertNotIn(reverse("attendance-regularisation-import"), html)
        self.assertNotIn("Upload filled sheet", html)
